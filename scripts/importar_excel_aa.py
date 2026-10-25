"""Importa el Excel de aire acondicionado del cliente (FORMATO AUTOMATIZADO INFORMES AIRE ACONDICIONADO.xlsm)
y escribe en stdout un SQL idempotente con:

  * clientes y sedes (10_BASE_DATO_EQUIPOS)
  * tipos de equipo normalizados (09_EQUIVALENCIAS_TIPO: el Excel tiene ~22 escrituras para 13 tipos)
  * inventario de equipos por sede (site_equipment)
  * plantilla de checklist v1 por tipo (08_CATALOGOS), solo si el tipo aún no tiene una vigente
  * catálogo de hallazgos (11_CATALOGO_HALLAZGOS), sin duplicar los ya cargados

Uso:
  python3 scripts/importar_excel_aa.py ruta/al/archivo.xlsm > /tmp/importacion.sql
  psql "<conexión>" -v ON_ERROR_STOP=1 -1 -f /tmp/importacion.sql

Volver a ejecutarlo actualiza los equipos existentes (por sede + código) sin tocar los servicios ya registrados.
"""
import re
import sys
import unicodedata
import warnings
from collections import Counter, OrderedDict

import openpyxl

warnings.filterwarnings('ignore')


def norm(s):
    s = unicodedata.normalize('NFKD', str(s or '')).encode('ascii', 'ignore').decode()
    return re.sub(r'\s+', ' ', s).strip().lower()


def clean(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = re.sub(r'\s+', ' ', str(v)).strip()
    return s or None


# Nombre canónico y unidad por defecto de cada tipo (según la hoja 09_EQUIVALENCIAS_TIPO)
CANON = {
    'cortina de aire': 'Cortina de aire', 'cortinas de aire': 'Cortina de aire',
    'fan coil agua helada': 'Fan coil agua helada', 'fan coil de agua helada': 'Fan coil agua helada',
    'fancoil de agua helada': 'Fan coil agua helada', 'fan coil chilled': 'Fan coil agua helada',
    'fancoil expansion directa': 'Fan coil expansión directa', 'fan coil expansion directa': 'Fan coil expansión directa',
    'fan coil de expansion directa': 'Fan coil expansión directa', 'fan coil xd': 'Fan coil expansión directa',
    'fan coil vrf': 'Fan coil VRF',
    'split ducto': 'Split ducto', 'split ducto (fan coil)': 'Split ducto',
    'split piso techo': 'Split piso techo', 'piso techo': 'Split piso techo',
    'split decorativo de pared': 'Split decorativo de pared', 'split decorativo': 'Split decorativo de pared',
    'split decorativo pared': 'Split decorativo de pared', 'split pared': 'Split decorativo de pared',
    'split decorativo piso techo': 'Split decorativo piso techo',
    'equipo autocontenido': 'Equipo autocontenido', 'autocontenido de agua helada': 'Equipo autocontenido',
    'autocontenido enfriado por agua': 'Equipo autocontenido', 'packaged (paquete 2 sistemas)': 'Equipo autocontenido',
    'centrifugos': 'Ventilador centrífugo', 'condensador fan coil': 'Condensador fan coil',
    'damper motorizado': 'Damper motorizado', 'extractores de aire': 'Extractor de aire',
}
UNIT = {'Cortina de aire': 'cm', 'Damper motorizado': 'W', 'Extractor de aire': 'CFM', 'Ventilador centrífugo': 'CFM'}

# Tipos sin rutina en el Excel: se les crea un checklist básico
GENERIC = [
    ('General', 'Inspección general del equipo'), ('Limpieza', 'Limpieza general del equipo'),
    ('Eléctrico', 'Revisión de conexiones eléctricas'), ('Eléctrico', 'Medición de tensión'),
    ('Eléctrico', 'Medición de corriente'), ('Operación', 'Prueba de funcionamiento'),
]

SYSTEM = {'a.a.': 'AA', 'aa': 'AA', 'vent.': 'VM', 'vm': 'VM', 'v.m.': 'VM', 'ext': 'EXT', 'dam': 'DAM', 'inyec.': 'INY'}


def canon(raw):
    c = CANON.get(norm(raw))
    if not c:
        raise SystemExit(f'Tipo de equipo sin equivalencia: {raw!r}')
    return c


def value_unit(activity):
    a = norm(activity)
    if 'tension' in a:
        return 'V'
    if 'corriente' in a:
        return 'A'
    if 'temperatura' in a or 'δt' in activity.lower():
        return '°C'
    if 'presion' in a:
        return 'PSI'
    return None


def q(v):
    return 'null' if v is None else "'" + str(v).replace("'", "''") + "'"


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = ['-- Generado por scripts/importar_excel_aa.py', 'set search_path = public;']

    # ── Inventario ──
    rows = [r for r in wb['10_BASE_DATO_EQUIPOS'].iter_rows(min_row=2, values_only=True) if any(r)]
    sites = OrderedDict()
    for r in rows:
        client, _prov, code, site, address, system, _qty, typ, cap, unit, brand, model, serial = r[:13]
        site = clean(site)
        # 6 filas vienen sin cliente: son tiendas Marathon de Superdeporte
        client = clean(client) or 'SUPERDEPORTE PLUS PERU'
        # Bajo Superdeporte figura "NIKE LARCOMAR" con códigos OFP8 y dirección del piso 8 de la torre: son oficinas
        if client == 'SUPERDEPORTE PLUS PERU' and site == 'NIKE LARCOMAR':
            site = 'OFICINAS PISO 8'
        s = sites.setdefault((client, site), {'address': Counter(), 'eq': []})
        if clean(address):
            s['address'][clean(address)] += 1
        unit = clean(unit)
        s['eq'].append({
            'code': clean(code), 'type': canon(typ), 'system': SYSTEM.get(norm(system), clean(system)),
            'capacity': None if clean(cap) in (None, '-') else clean(cap),
            'unit': None if unit in (None, '-') else ('cm' if unit.lower() == 'cm' else unit.upper()),
            'brand': clean(brand), 'model': clean(model), 'serial': clean(serial),
        })

    types = sorted({e['type'] for s in sites.values() for e in s['eq']} | set(CANON.values()))
    out.append('\n-- Tipos de equipo')
    out += [f"insert into equipment_types (name, default_unit) values ({q(t)}, {q(UNIT.get(t, 'BTU'))}) on conflict (name) do nothing;" for t in types]

    out.append('\n-- Clientes y sedes')
    for c in sorted({c for c, _ in sites}):
        out.append(f'insert into clients (name) values ({q(c)}) on conflict (name) do nothing;')
    for (c, name), s in sites.items():
        address = s['address'].most_common(1)[0][0] if s['address'] else None
        prefixes = Counter(re.sub(r'[\s\-]*\d+$', '', e['code']) for e in s['eq'])
        prefix = prefixes.most_common(1)[0][0]
        out.append(
            f"insert into sites (client_id, name, address, equipment_prefix) "
            f"select id, {q(name)}, {q(address)}, {q(prefix)} from clients where name = {q(c)} "
            f"on conflict (client_id, name) do update set address = coalesce(excluded.address, sites.address), "
            f"equipment_prefix = coalesce(sites.equipment_prefix, excluded.equipment_prefix);")

    out.append('\n-- Inventario de equipos por sede')
    for (c, name), s in sites.items():
        for i, e in enumerate(s['eq'], 1):
            out.append(
                'insert into site_equipment (site_id, code, equipment_type_id, system, capacity, unit, brand, model, serial, position) '
                f"select st.id, {q(e['code'])}, et.id, {q(e['system'])}, {q(e['capacity'])}, {q(e['unit'])}, {q(e['brand'])}, "
                f"{q(e['model'])}, {q(e['serial'])}, {i} "
                f"from sites st join clients cl on cl.id = st.client_id, equipment_types et "
                f"where cl.name = {q(c)} and st.name = {q(name)} and et.name = {q(e['type'])} "
                'on conflict (site_id, code) do update set equipment_type_id = excluded.equipment_type_id, system = excluded.system, '
                'capacity = excluded.capacity, unit = excluded.unit, brand = excluded.brand, model = excluded.model, '
                'serial = excluded.serial, position = excluded.position;')

    # ── Checklist por tipo ──
    routines = OrderedDict()
    for r in wb['08_CATALOGOS'].iter_rows(min_row=2, values_only=True):
        if not r[0] or norm(r[4]) not in ('si', 'sí'):
            continue
        t = canon(r[0])
        # Varias escrituras del mismo tipo repiten la rutina: se toma la primera
        first = routines.setdefault(t, {'raw': r[0], 'items': []})
        if first['raw'] == r[0]:
            first['items'].append((clean(r[3]) or 'General', clean(r[2])))
    out.append('\n-- Plantillas de checklist (v1, solo para tipos sin plantilla vigente)')
    for t in types:
        items = routines[t]['items'] if t in routines else GENERIC
        values = ', '.join(
            f"({i}, {q(cat)}, {q(act)}, {q('cn_value' if value_unit(act) else 'cn')}, {q(value_unit(act))})"
            for i, (cat, act) in enumerate(items, 1))
        out.append(f"""
with t as (select id from equipment_types where name = {q(t)}
           and not exists (select 1 from checklist_templates ct where ct.equipment_type_id = equipment_types.id and ct.is_current)),
     tpl as (insert into checklist_templates (equipment_type_id, version, is_current) select id, 1, true from t returning id)
insert into checklist_template_items (template_id, position, category, activity, response_type, value_unit, required)
select tpl.id, v.pos, v.cat, v.act, v.rt, v.unit, true from tpl,
  (values {values}) as v(pos, cat, act, rt, unit);""")

    # ── Catálogo de hallazgos ──
    seen = set()
    findings = []
    for r in wb['11_CATALOGO_HALLAZGOS'].iter_rows(min_row=2, values_only=True):
        if not r[0] or not r[1]:
            continue
        t = canon(r[0])
        key = (t, norm(r[1]))
        if key in seen:
            continue
        seen.add(key)
        prio = clean(r[5]) if clean(r[5]) in ('Alta', 'Media', 'Baja') else 'Media'
        findings.append((t, clean(r[1]), clean(r[2]), clean(r[3]), clean(r[4]), prio))
    out.append('\n-- Catálogo de hallazgos (códigos H-### a continuación de los existentes)')
    out.append('create temp table _hallazgos (n serial, tipo text, finding text, cause text, action text, rec text, prio text) on commit drop;')
    out.append('insert into _hallazgos (tipo, finding, cause, action, rec, prio) values\n' + ',\n'.join(
        f'({q(t)}, {q(f)}, {q(c)}, {q(a)}, {q(rc)}, {q(p)})' for t, f, c, a, rc, p in findings) + ';')
    out.append("""
insert into findings_catalog (code, finding, equipment_type_id, probable_cause, suggested_action, recommendation, priority)
select 'H-' || lpad((base.max + row_number() over (order by h.n))::text, 3, '0'), h.finding, et.id, h.cause, h.action, h.rec, h.prio
from _hallazgos h
join equipment_types et on et.name = h.tipo
cross join (select coalesce(max(nullif(regexp_replace(code, '\\D', '', 'g'), '')::int), 0) as max from findings_catalog) base
where not exists (
  select 1 from findings_catalog fc where fc.equipment_type_id = et.id and lower(fc.finding) = lower(h.finding)
);""")

    total = sum(len(s['eq']) for s in sites.values())
    print('\n'.join(out))
    print(f'-- {len(sites)} sedes · {total} equipos · {len(types)} tipos · {len(findings)} hallazgos', file=sys.stderr)


if __name__ == '__main__':
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    main(sys.argv[1])
